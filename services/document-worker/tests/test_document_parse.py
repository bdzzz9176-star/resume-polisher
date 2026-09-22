from io import BytesIO

import pymupdf
from docx import Document
from fastapi.testclient import TestClient

from app.main import app


client = TestClient(app)


def make_docx() -> bytes:
    stream = BytesIO()
    document = Document()
    title = document.add_paragraph()
    run = title.add_run("张三 - 嵌入式软件工程师")
    run.bold = True
    document.add_paragraph("项目经历：使用摄像头完成图像采集。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "技能"
    table.cell(0, 1).text = "C / Linux"
    document.save(stream)
    return stream.getvalue()


def make_text_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Embedded camera project with Linux and image capture experience.")
    content = document.tobytes()
    document.close()
    return content


def make_empty_pdf() -> bytes:
    document = pymupdf.open()
    document.new_page()
    content = document.tobytes()
    document.close()
    return content


def test_parse_docx_preserves_block_order() -> None:
    response = client.post(
        "/v1/documents/parse",
        files={"file": ("resume.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert response.status_code == 200
    result = response.json()
    assert result["contract"] == "document-parse.v1"
    assert result["documentKind"] == "docx"
    assert [block["kind"] for block in result["blocks"]] == ["paragraph", "paragraph", "table"]
    assert result["blocks"][0]["runs"][0]["bold"] is True
    assert "摄像头" in result["plainText"]


def test_parse_text_pdf_returns_positioned_blocks() -> None:
    response = client.post(
        "/v1/documents/parse",
        files={"file": ("resume.pdf", make_text_pdf(), "application/pdf")},
    )

    assert response.status_code == 200
    result = response.json()
    assert result["documentKind"] == "pdf"
    assert result["blocks"][0]["page"] == 1
    assert result["blocks"][0]["bbox"] is not None
    assert result["warnings"] == []


def test_sparse_pdf_returns_scan_warning() -> None:
    response = client.post(
        "/v1/documents/parse",
        files={"file": ("scan.pdf", make_empty_pdf(), "application/pdf")},
    )

    assert response.status_code == 200
    assert response.json()["warnings"][0]["code"] == "PDF_TEXT_TOO_SPARSE"


def test_rejects_unsupported_file_type() -> None:
    response = client.post(
        "/v1/documents/parse",
        files={"file": ("resume.txt", b"hello", "text/plain")},
    )

    assert response.status_code == 400
    assert response.json()["detail"]["code"] == "UNSUPPORTED_FILE_TYPE"

