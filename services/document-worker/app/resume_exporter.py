from io import BytesIO
from pathlib import Path
import base64
from collections import Counter

import fitz
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Cm, Pt, RGBColor

from app.models import ResumeExportRequest


def _run_format(run) -> dict:
    return {
        "bold": run.bold,
        "italic": run.italic,
        "underline": run.underline,
        "font_name": run.font.name,
        "font_size": run.font.size,
        "color": run.font.color.rgb if run.font.color and run.font.color.rgb else None,
    }


def _apply_run_format(run, sample: dict) -> None:
    if sample.get("bold") is not None:
        run.bold = sample["bold"]
    if sample.get("italic") is not None:
        run.italic = sample["italic"]
    if sample.get("underline") is not None:
        run.underline = sample["underline"]
    if sample.get("font_name"):
        run.font.name = sample["font_name"]
    if sample.get("font_size"):
        run.font.size = sample["font_size"]
    if sample.get("color"):
        run.font.color.rgb = sample["color"]


def _style_exists(document: Document, style_name: str | None) -> bool:
    if not style_name:
        return False
    return any(style.name == style_name for style in document.styles)


def _sample_format(paragraph) -> dict:
    if paragraph and paragraph.runs:
        run = next((candidate for candidate in paragraph.runs if candidate.text.strip()), paragraph.runs[0])
        return _run_format(run)
    return {}


def _paragraph_samples(document: Document) -> dict:
    paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
    first = paragraphs[0] if paragraphs else None
    heading = next(
        (
            paragraph
            for paragraph in paragraphs[1:]
            if (paragraph.style and "heading" in paragraph.style.name.lower())
            or any(run.bold for run in paragraph.runs)
        ),
        first,
    )
    body_paragraph = next((paragraph for paragraph in paragraphs[1:] if paragraph.runs), first)
    style_names = [
        paragraph.style.name
        for paragraph in paragraphs
        if paragraph.style and paragraph.style.name and paragraph.text.strip()
    ]
    body_style = Counter(style_names).most_common(1)[0][0] if style_names else "Normal"
    list_style = next(
        (
            paragraph.style.name
            for paragraph in paragraphs
            if paragraph.style and "list" in paragraph.style.name.lower()
        ),
        body_style,
    )

    return {
        "title_style": first.style.name if first and first.style else "Normal",
        "heading_style": heading.style.name if heading and heading.style else body_style,
        "body_style": body_style,
        "list_style": list_style,
        "title_format": _sample_format(first),
        "heading_format": _sample_format(heading),
        "body_format": _sample_format(body_paragraph),
    }


def _clear_document_body(document: Document) -> None:
    body = document._body._element  # noqa: SLF001 - python-docx has no public body clearing API.
    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)


def _add_formatted_paragraph(document: Document, text: str, style_name: str, sample: dict):
    style = style_name if _style_exists(document, style_name) else "Normal"
    paragraph = document.add_paragraph(style=style)
    run = paragraph.add_run(text)
    _apply_run_format(run, sample)
    return paragraph


def export_docx_from_template(resume: ResumeExportRequest) -> bytes:
    if not resume.source_docx_base64:
        return export_docx(resume)

    source = base64.b64decode(resume.source_docx_base64)
    document = Document(BytesIO(source))
    samples = _paragraph_samples(document)
    _clear_document_body(document)

    title = _add_formatted_paragraph(
        document,
        resume.title.strip() or "优化简历",
        samples["title_style"],
        samples["title_format"],
    )
    title.paragraph_format.space_after = Pt(5)

    if resume.summary.strip():
        summary = _add_formatted_paragraph(
            document,
            resume.summary.strip(),
            samples["body_style"],
            samples["body_format"],
        )
        summary.paragraph_format.space_after = Pt(8)

    for export_section in resume.sections:
        heading = _add_formatted_paragraph(
            document,
            export_section.heading.strip(),
            samples["heading_style"],
            samples["heading_format"],
        )
        heading.paragraph_format.space_before = Pt(6)
        heading.paragraph_format.space_after = Pt(3)

        for item in export_section.items:
            paragraph = _add_formatted_paragraph(
                document,
                item.strip(),
                samples["list_style"],
                samples["body_format"],
            )
            paragraph.paragraph_format.space_after = Pt(2)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def export_docx(resume: ResumeExportRequest) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    title_run = title.add_run(resume.title)
    title_run.bold = True
    title_run.font.name = "Microsoft YaHei"
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(0x10, 0x20, 0x2A)

    if resume.summary.strip():
        summary = document.add_paragraph(resume.summary.strip())
        summary.paragraph_format.space_after = Pt(8)
        for run in summary.runs:
            run.font.color.rgb = RGBColor(0x4D, 0x61, 0x6B)

    for export_section in resume.sections:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(7)
        heading.paragraph_format.space_after = Pt(4)
        heading_run = heading.add_run(export_section.heading)
        heading_run.bold = True
        heading_run.font.name = "Microsoft YaHei"
        heading_run.font.size = Pt(12.5)
        heading_run.font.color.rgb = RGBColor(0x00, 0x8F, 0x95)

        for item in export_section.items:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Cm(0.45)
            paragraph.paragraph_format.first_line_indent = Cm(-0.2)
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(item.strip())
            run.font.name = "Microsoft YaHei"
            run.font.size = Pt(10.5)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _font_path() -> str:
    candidates = [
        Path("C:/Windows/Fonts/simhei.ttf"),
        Path("C:/Windows/Fonts/msyh.ttc"),
        Path("C:/Windows/Fonts/msyhbd.ttc"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return str(candidate)
    raise RuntimeError("No Chinese font was found for PDF export.")


def _wrap_text(text: str, max_chars: int) -> list[str]:
    lines: list[str] = []
    for raw_line in text.splitlines() or [""]:
        current = raw_line.strip()
        if not current:
            lines.append("")
            continue
        while len(current) > max_chars:
            lines.append(current[:max_chars])
            current = current[max_chars:]
        if current:
            lines.append(current)
    return lines


def export_pdf(resume: ResumeExportRequest) -> bytes:
    document = fitz.open()
    font_path = _font_path()
    page = document.new_page(width=595, height=842)
    font_name = "resume-font"
    page.insert_font(fontname=font_name, fontfile=font_path)
    y = 58.0

    def new_page() -> None:
        nonlocal page, y
        page = document.new_page(width=595, height=842)
        page.insert_font(fontname=font_name, fontfile=font_path)
        y = 54.0

    def write_lines(text: str, size: float, color: tuple[float, float, float], indent: float = 0) -> None:
        nonlocal y
        max_chars = 42 if size >= 15 else 56
        line_height = size * 1.65
        for line in _wrap_text(text, max_chars):
            if y + line_height > 800:
                new_page()
            page.insert_text(
                (54 + indent, y),
                line,
                fontsize=size,
                fontname=font_name,
                color=color,
            )
            y += line_height

    write_lines(resume.title, 20, (0.06, 0.13, 0.17))
    y += 2
    if resume.summary.strip():
        write_lines(resume.summary.strip(), 10.5, (0.30, 0.38, 0.42))
        y += 8

    for export_section in resume.sections:
        if y > 760:
            new_page()
        write_lines(export_section.heading, 13, (0.0, 0.56, 0.58))
        y += 2
        for item in export_section.items:
            write_lines(f"- {item.strip()}", 10.5, (0.10, 0.16, 0.19), indent=8)
            y += 2
        y += 5

    output = document.tobytes(garbage=4, deflate=True)
    document.close()
    return output

